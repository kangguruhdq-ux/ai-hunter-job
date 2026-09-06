import os
from typing import Tuple, Optional
from pypdf import PdfReader
from docx import Document
from app.core.logging import logger

class ResumeParsingError(Exception):
    """Custom exception raised when resume text extraction fails."""
    pass

class ResumeAgent:
    """
    Agent responsible for extracting clean text from uploaded resumes (PDF, DOCX).
    Ensures safe file handling, detects corrupted files, and strips artifacts.
    Preserves sequential reading order of paragraphs, headings, and tables.
    """

    ALLOWED_EXTENSIONS = {".pdf", ".docx"}

    @classmethod
    def validate_file(cls, file_name: str, file_size: int, max_size_bytes: int = 10 * 1024 * 1024) -> str:
        """Validates file extension and size. Returns lowercased extension."""
        if not file_name:
            raise ResumeParsingError("No file name provided.")

        _, ext = os.path.splitext(file_name.lower())
        if ext not in cls.ALLOWED_EXTENSIONS:
            raise ResumeParsingError(
                f"Unsupported file format '{ext}'. Only PDF and DOCX files are supported."
            )

        if file_size == 0:
            raise ResumeParsingError("The uploaded file is empty (0 bytes).")

        if file_size > max_size_bytes:
            max_mb = max_size_bytes / (1024 * 1024)
            raise ResumeParsingError(
                f"File size exceeds maximum allowed limit of {max_mb:.0f} MB."
            )

        return ext

    @classmethod
    def extract_text(cls, file_path: str, ext: str) -> Tuple[str, Optional[int]]:
        """
        Extracts raw text and page count from the specified file.
        Returns: (cleaned_text, page_count)
        """
        if not os.path.exists(file_path):
            raise ResumeParsingError(f"File not found at path: {file_path}")

        ext = ext.lower().strip()
        if not ext.startswith("."):
            ext = f".{ext}"

        if ext == ".pdf":
            return cls._extract_from_pdf(file_path)
        elif ext == ".docx":
            return cls._extract_from_docx(file_path)
        else:
            raise ResumeParsingError(f"Cannot parse file type: {ext}")

    @classmethod
    def _extract_from_pdf(cls, file_path: str) -> Tuple[str, int]:
        try:
            reader = PdfReader(file_path)
            page_count = len(reader.pages)
            if page_count == 0:
                raise ResumeParsingError("PDF file contains no pages.")

            extracted_parts = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    extracted_parts.append(text.strip())

            full_text = "\n\n".join(extracted_parts).strip()
            if not full_text or len(full_text.strip()) < 10:
                raise ResumeParsingError(
                    "The PDF file contains no extractable text. It may be a scanned image or empty."
                )

            return full_text, page_count
        except ResumeParsingError:
            raise
        except Exception as e:
            logger.error(f"Error reading PDF file {file_path}: {e}")
            raise ResumeParsingError(
                "Unable to read this PDF resume. The file may be password-protected or corrupted."
            )

    @classmethod
    def _extract_from_docx(cls, file_path: str) -> Tuple[str, Optional[int]]:
        try:
            doc = Document(file_path)
            extracted_elements = []

            # Traverse body elements in visual document order to preserve context
            if hasattr(doc, "element") and hasattr(doc.element, "body"):
                for child in doc.element.body:
                    if child.tag.endswith("p"):
                        p_text = "".join(child.itertext()).strip()
                        if p_text:
                            extracted_elements.append(p_text)
                    elif child.tag.endswith("tbl"):
                        for row_elem in child.findall(".//{*}tr"):
                            row_cells = []
                            for cell_elem in row_elem.findall(".//{*}tc"):
                                cell_text = "".join(cell_elem.itertext()).strip()
                                if cell_text and (not row_cells or cell_text != row_cells[-1]):
                                    row_cells.append(cell_text)
                            if row_cells:
                                extracted_elements.append(" | ".join(row_cells))

            # Fallback if body iteration was empty
            if not extracted_elements:
                for p in doc.paragraphs:
                    if p.text.strip():
                        extracted_elements.append(p.text.strip())
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                        if row_text:
                            extracted_elements.append(row_text)

            full_text = "\n\n".join(extracted_elements).strip()
            if not full_text or len(full_text.strip()) < 10:
                raise ResumeParsingError(
                    "The DOCX file contains no extractable text or is empty."
                )

            return full_text, None
        except ResumeParsingError:
            raise
        except Exception as e:
            logger.error(f"Error reading DOCX file {file_path}: {e}")
            raise ResumeParsingError(
                "Unable to read this DOCX file. The document may be corrupted."
            )
