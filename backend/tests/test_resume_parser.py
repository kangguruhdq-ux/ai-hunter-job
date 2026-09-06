import io
import os
import tempfile
import pytest
from docx import Document
from pypdf import PdfWriter
from fastapi.testclient import TestClient

from app.main import app
from app.agents.resume_agent import ResumeAgent, ResumeParsingError
from app.core.database import Base, engine, SessionLocal

@pytest.fixture(scope="module", autouse=True)
def setup_test_db():
    Base.metadata.create_all(bind=engine)
    yield


@pytest.fixture
def client():
    return TestClient(app)

def create_sample_docx(text: str) -> str:
    temp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    doc = Document()
    doc.add_heading("Alex Mercer - Senior Software Engineer", level=1)
    doc.add_paragraph(text)
    doc.save(temp.name)
    temp.close()
    return temp.name

def test_resume_agent_validation():
    # Unsupported extension
    with pytest.raises(ResumeParsingError, match="Unsupported file format"):
        ResumeAgent.validate_file("resume.exe", 1000)

    # Empty file
    with pytest.raises(ResumeParsingError, match="The uploaded file is empty"):
        ResumeAgent.validate_file("resume.pdf", 0)

    # Oversized file
    with pytest.raises(ResumeParsingError, match="File size exceeds"):
        ResumeAgent.validate_file("resume.pdf", 20 * 1024 * 1024, max_size_bytes=10 * 1024 * 1024)

    # Valid PDF
    ext = ResumeAgent.validate_file("my_cv.PDF", 5000)
    assert ext == ".pdf"

def test_docx_text_extraction():
    content = (
        "Experienced Backend Developer with 5+ years of experience in Python, FastAPI, "
        "PostgreSQL, Docker, and Kubernetes. Built high-scale microservices handling 50k req/sec."
    )
    docx_path = create_sample_docx(content)
    try:
        extracted, page_count = ResumeAgent.extract_text(docx_path, ".docx")
        assert "Alex Mercer" in extracted
        assert "FastAPI" in extracted
        assert "PostgreSQL" in extracted
    finally:
        if os.path.exists(docx_path):
            os.unlink(docx_path)

def test_corrupted_file_handling():
    temp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
    temp.write(b"not a valid pdf content at all corrupted garbage")
    temp.close()
    try:
        with pytest.raises(ResumeParsingError):
            ResumeAgent.extract_text(temp.name, ".pdf")
    finally:
        if os.path.exists(temp.name):
            os.unlink(temp.name)

def test_resume_upload_api_docx(client):
    content = (
        "Full Stack Engineer with 6 years building modern web apps. "
        "Proficient in Next.js, React, TypeScript, Python, and AWS."
    )
    docx_path = create_sample_docx(content)
    try:
        with open(docx_path, "rb") as f:
            response = client.post(
                "/api/v1/resume/upload",
                files={"file": ("test_resume.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            )
        assert response.status_code == 201
        data = response.json()
        assert data["file_name"] == "test_resume.docx"
        assert data["file_type"] == "docx"
        assert data["status"] == "parsed"

        # Check retrieve latest
        latest_res = client.get("/api/v1/resume/latest")
        assert latest_res.status_code == 200
        latest_data = latest_res.json()
        assert latest_data["id"] == data["id"]
        assert "Full Stack Engineer" in latest_data["raw_text"]
    finally:
        if os.path.exists(docx_path):
            os.unlink(docx_path)

def test_resume_upload_api_invalid_file(client):
    response = client.post(
        "/api/v1/resume/upload",
        files={"file": ("bad_file.txt", b"plain text is not supported", "text/plain")}
    )
    assert response.status_code == 400
    assert "Unsupported file format" in response.json()["detail"]
