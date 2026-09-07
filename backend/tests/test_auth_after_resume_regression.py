import io
import pytest
from docx import Document
from fastapi.testclient import TestClient
from app.main import app
from app.core.database import SessionLocal
from app.models.user import User

@pytest.fixture(scope="module")
def client():
    return TestClient(app)

def create_valid_docx_bytes(title: str = "Mahabbah Mahabban Romadhon"):
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph("Teknik Komputer dan Jaringan SMKN 3 YOGYAKARTA")
    doc.add_paragraph("Pengalaman PKL di Immersa. Troubleshooting jaringan dan konfigurasi router MikroTik.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

def test_login_flow_stability_after_resume_processing_sequence(client):
    email = "regression_resume_auth@jobhunter.ai"
    password = "SecurePassword123!"

    with SessionLocal() as db:
        db.query(User).filter(User.email == email).delete(synchronize_session=False)
        db.commit()

    # 1. Register
    reg_res = client.post("/api/v1/auth/register", json={
        "email": email,
        "full_name": "Regression Tester",
        "password": password
    })
    assert reg_res.status_code == 201, f"Registration failed: {reg_res.text}"
    token_1 = reg_res.json()["access_token"]
    headers_1 = {"Authorization": f"Bearer {token_1}"}

    # 2. Login
    login_res = client.post("/api/v1/auth/login", json={
        "email": email,
        "password": password
    })
    assert login_res.status_code == 200, f"Login failed: {login_res.text}"
    token_active = login_res.json()["access_token"]
    headers_active = {"Authorization": f"Bearer {token_active}"}

    # 3. Open dashboard stats
    dash_res = client.get("/api/v1/ai/dashboard-stats", headers=headers_active)
    assert dash_res.status_code == 200

    # 4. Upload resume
    docx_bytes = create_valid_docx_bytes("Mahabbah Mahabban Romadhon")
    upload_res = client.post(
        "/api/v1/resume/upload",
        files={"file": ("resume_1.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=headers_active
    )
    assert upload_res.status_code == 201, f"Resume upload failed: {upload_res.text}"
    resume_id = upload_res.json()["id"]

    # 5. Run AI analysis
    analyze_res = client.post(
        f"/api/v1/candidate/profile/extract/{resume_id}",
        headers=headers_active
    )
    assert analyze_res.status_code == 200, f"Profile extract failed: {analyze_res.text}"
    profile_data = analyze_res.json()
    assert profile_data["name"] is not None

    # 6. Logout
    logout_res = client.post("/api/v1/auth/logout", headers=headers_active)
    assert logout_res.status_code == 200

    # 7. Login kembali (CRITICAL: must succeed without backend restart!)
    login_again_res = client.post("/api/v1/auth/login", json={
        "email": email,
        "password": password
    })
    assert login_again_res.status_code == 200, f"Login after resume processing failed: {login_again_res.text}"
    token_2 = login_again_res.json()["access_token"]
    headers_2 = {"Authorization": f"Bearer {token_2}"}

    # 8. Verify session with /auth/me
    me_res = client.get("/api/v1/auth/me", headers=headers_2)
    assert me_res.status_code == 200
    assert me_res.json()["email"] == email

    # 9. Navigate dashboard
    dash_res2 = client.get("/api/v1/ai/dashboard-stats", headers=headers_2)
    assert dash_res2.status_code == 200

    # 10. Upload another resume & run analysis
    upload_res2 = client.post(
        "/api/v1/resume/upload",
        files={"file": ("resume_2.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=headers_2
    )
    assert upload_res2.status_code == 201
    resume_id2 = upload_res2.json()["id"]

    analyze_res2 = client.post(
        f"/api/v1/candidate/profile/extract/{resume_id2}",
        headers=headers_2
    )
    assert analyze_res2.status_code == 200

    # 11. Logout
    logout_res2 = client.post("/api/v1/auth/logout", headers=headers_2)
    assert logout_res2.status_code == 200

    # 12. Login again
    final_login = client.post("/api/v1/auth/login", json={
        "email": email,
        "password": password
    })
    assert final_login.status_code == 200
    assert "access_token" in final_login.json()
